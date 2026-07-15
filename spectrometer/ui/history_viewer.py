"""最多 12 页、每页最多 64 条曲线的历史查看器。"""

import csv
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from ..qt import QtWidgets
from .plot_widget import HistoryPlotWidget


class HistoryViewer(QtWidgets.QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        layout = QtWidgets.QVBoxLayout(self); layout.setContentsMargins(0, 0, 0, 0)
        toolbar = QtWidgets.QHBoxLayout()
        open_button = QtWidgets.QPushButton("打开 CSV / Excel")
        open_button.clicked.connect(self.open_files); toolbar.addWidget(open_button)
        new_button = QtWidgets.QPushButton("新建页面"); new_button.clicked.connect(self.new_page); toolbar.addWidget(new_button)
        clear_button = QtWidgets.QPushButton("清空当前页"); clear_button.clicked.connect(self.clear_current); toolbar.addWidget(clear_button)
        toolbar.addStretch(1); layout.addLayout(toolbar)
        self.tabs = QtWidgets.QTabWidget(); self.tabs.setTabsClosable(True)
        self.tabs.tabCloseRequested.connect(self._close_tab); layout.addWidget(self.tabs, 1)
        self.new_page("历史 1")

    def new_page(self, title=None):
        if self.tabs.count() >= 12:
            return None
        plot = HistoryPlotWidget(); index = self.tabs.addTab(plot, title or f"历史 {self.tabs.count() + 1}")
        self.tabs.setCurrentIndex(index); return plot

    def current_plot(self):
        return self.tabs.currentWidget()

    def clear_current(self):
        plot = self.current_plot()
        if plot: plot.clear()

    def _close_tab(self, index):
        if self.tabs.count() <= 1:
            self.clear_current(); return
        widget = self.tabs.widget(index); self.tabs.removeTab(index); widget.deleteLater()

    def open_files(self):
        paths, _ = QtWidgets.QFileDialog.getOpenFileNames(
            self, "打开光谱数据", "", "光谱数据 (*.csv *.xlsx)"
        )
        for path in paths:
            self.load_file(path)

    def load_file(self, path):
        path = Path(path)
        plot = self.current_plot() or self.new_page(path.stem)
        if path.suffix.lower() == ".csv":
            self._load_csv(path, plot)
        elif path.suffix.lower() == ".xlsx":
            self._load_xlsx(path, plot)
        else:
            raise ValueError("仅支持 CSV 和 Excel")

    @staticmethod
    def _load_csv(path, plot):
        with path.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file))
        header_index = next((index for index, row in enumerate(rows) if row and row[0] == "Pixel"), None)
        if header_index is None:
            raise ValueError("CSV 中未找到 Pixel 表头")
        header = rows[header_index]; data = rows[header_index + 1:]
        x_column = 1 if len(header) > 1 and header[1] == "Wavelength" else 0
        x = np.array([float(row[x_column] or row[0]) for row in data if row and len(row) > x_column])
        for column in range(2, len(header)):
            values = []
            for row in data[:len(x)]:
                values.append(float(row[column]) if len(row) > column and row[column] else np.nan)
            plot.load_spectrum(x, np.asarray(values), header[column])

    @staticmethod
    def _load_xlsx(path, plot):
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            if sheet.title == "采集概要": continue
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header or header[0] != "Pixel": continue
            data = list(rows); x = np.array([row[1] if row[1] is not None else row[0] for row in data], dtype=float)
            for column in range(2, len(header)):
                values = np.array([row[column] if len(row) > column and row[column] is not None else np.nan for row in data], dtype=float)
                plot.load_spectrum(x, values, f"{sheet.title}/{header[column]}")
