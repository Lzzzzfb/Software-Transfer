import os
import threading
import time

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook

from spectrometer.device.spectrometer import SpectrometerDevice
from spectrometer.domain.enums import (
    AcquisitionMode,
    AcquisitionOwner,
    StorageFormat,
    SyncMode,
)
from spectrometer.domain.models import AcquisitionRequest, SpectrumFrame
from spectrometer.qt import QtCore
from spectrometer.storage import session_manager
from spectrometer.storage.session_manager import StorageSessionManager


_APPLICATION = None


def application():
    global _APPLICATION
    _APPLICATION = QtCore.QCoreApplication.instance() or QtCore.QCoreApplication([])
    return _APPLICATION


class FakeCoordinator:
    instances = []
    close_gate = None
    close_error = None

    def __init__(self, output_directory, session, **kwargs):
        self.output_directory = output_directory
        self.session = session
        self.frames = []
        self.errors = []
        self.exported_files = []
        self.started = False
        self.closed = False
        FakeCoordinator.instances.append(self)

    @property
    def queue_ratio(self):
        return 0.0

    def start(self):
        self.started = True

    def submit(self, frame):
        self.frames.append(frame)

    def close(self):
        if self.close_gate is not None:
            self.close_gate.wait(1)
        if self.close_error is not None:
            raise self.close_error
        self.closed = True


def ready_device(device_id):
    device = SpectrometerDevice(device_id, f"COM{device_id + 10}")
    device.connected = True
    device.initialized = True
    device.info.prod_serial = f"00{device_id + 1}"
    device.info.valid_pixel = 2
    device.info.pixel_count = 2
    return device


def request(device_id):
    return AcquisitionRequest.create(
        AcquisitionOwner.LOCAL,
        [device_id],
        auto_store=True,
    )


def spectrum(device_id, sequence):
    return SpectrumFrame.create(
        device_id,
        sequence << 8,
        [sequence, sequence + 1],
        monotonic_ns=sequence,
        timestamp_ns=sequence,
    )


def wait_until(predicate, timeout=1.0):
    app = application()
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        app.processEvents()
        if predicate():
            return True
        time.sleep(0.005)
    return predicate()


def test_two_local_sessions_route_and_close_independently(tmp_path):
    FakeCoordinator.instances = []
    FakeCoordinator.close_gate = None
    FakeCoordinator.close_error = None
    manager = StorageSessionManager(
        tmp_path, coordinator_factory=FakeCoordinator
    )
    first = request(0)
    second = request(1)
    manager.start_session(first, [ready_device(0)])
    manager.start_session(second, [ready_device(1)])

    manager.submit(spectrum(0, 1))
    manager.submit(spectrum(1, 2))
    assert [frame.device_id for frame in FakeCoordinator.instances[0].frames] == [0]
    assert [frame.device_id for frame in FakeCoordinator.instances[1].frames] == [1]

    closed = []
    manager.session_closed.connect(lambda task_id, files, errors: closed.append(task_id))
    manager.close_session(first.task_id)
    assert wait_until(lambda: first.task_id in closed)
    assert manager.has_session(second.task_id)
    manager.submit(spectrum(1, 3))
    assert len(FakeCoordinator.instances[1].frames) == 2
    manager.close_session(second.task_id)
    assert wait_until(lambda: second.task_id in closed)


def test_slow_close_keeps_qt_event_loop_responsive(tmp_path):
    app = application()
    FakeCoordinator.instances = []
    FakeCoordinator.close_error = None
    gate = threading.Event()
    FakeCoordinator.close_gate = gate
    manager = StorageSessionManager(
        tmp_path, coordinator_factory=FakeCoordinator
    )
    acquisition = request(0)
    manager.start_session(acquisition, [ready_device(0)])
    ticks = []
    timer = QtCore.QTimer()
    timer.setTimerType(QtCore.Qt.PreciseTimer)
    timer.timeout.connect(lambda: ticks.append(time.monotonic()))
    timer.start(5)

    manager.close_session(acquisition.task_id)
    deadline = time.monotonic() + 0.08
    while time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.002)
    gate.set()

    assert len(ticks) >= 5
    assert wait_until(lambda: not manager.has_session(acquisition.task_id))
    timer.stop()
    FakeCoordinator.close_gate = None


def test_detached_session_cannot_receive_tail_frames(tmp_path):
    FakeCoordinator.instances = []
    FakeCoordinator.close_error = None
    gate = threading.Event()
    FakeCoordinator.close_gate = gate
    manager = StorageSessionManager(
        tmp_path, coordinator_factory=FakeCoordinator
    )
    acquisition = request(0)
    manager.start_session(acquisition, [ready_device(0)])
    manager.close_session(acquisition.task_id)

    assert not manager.submit(spectrum(0, 9))
    assert FakeCoordinator.instances[0].frames == []
    gate.set()
    assert wait_until(lambda: not manager.has_session(acquisition.task_id))
    FakeCoordinator.close_gate = None


def test_close_error_releases_session_and_is_reported(tmp_path):
    FakeCoordinator.instances = []
    FakeCoordinator.close_gate = None
    FakeCoordinator.close_error = RuntimeError("injected close failure")
    manager = StorageSessionManager(
        tmp_path, coordinator_factory=FakeCoordinator
    )
    acquisition = request(0)
    manager.start_session(acquisition, [ready_device(0)])
    closed = []
    manager.session_closed.connect(
        lambda task_id, files, errors: closed.append((task_id, errors))
    )

    assert manager.close_session(acquisition.task_id)
    assert wait_until(lambda: bool(closed))
    assert not manager.has_session(acquisition.task_id)
    assert closed == [(acquisition.task_id, ["injected close failure"])]
    FakeCoordinator.close_error = None


def test_real_coordinator_exports_human_named_multi_device_files(tmp_path):
    manager = StorageSessionManager(tmp_path)
    acquisition = AcquisitionRequest(
        task_id="integration-task",
        started_at="2026-07-22T15:30:45+08:00",
        owner=AcquisitionOwner.GLOBAL,
        device_ids=(0, 1),
        mode=AcquisitionMode.CONTINUOUS,
        sync_mode=SyncMode.SOFTWARE,
        auto_store=True,
        storage_format=StorageFormat.CSV_EXCEL,
        batch_size=2,
    )
    devices = [ready_device(0), ready_device(1)]
    manager.start_session(acquisition, devices)
    for sequence in (1, 2):
        for device_id in (0, 1):
            assert manager.submit(spectrum(device_id, sequence))

    closed = []
    manager.session_closed.connect(
        lambda task_id, files, errors: closed.append((task_id, files, errors))
    )
    assert manager.close_session(acquisition.task_id)
    assert wait_until(lambda: bool(closed), timeout=5.0)
    assert closed[0][0] == acquisition.task_id
    assert closed[0][2] == []

    names = {path.name for path in tmp_path.iterdir()}
    assert names == {
        "20260722_SN001_B0001.csv",
        "20260722_SN002_B0001.csv",
        "20260722_多设备_B0001.xlsx",
    }
    workbook_path = tmp_path / "20260722_多设备_B0001.xlsx"
    workbook = load_workbook(workbook_path, read_only=True)
    assert workbook.sheetnames == ["采集概要", "001", "002"]


def test_sealed_export_reports_outputs_recovery_files_and_cleanup_warnings(
    tmp_path, monkeypatch
):
    output = tmp_path / "batch.csv"
    recovery = tmp_path / "session.zgs"
    recovery.write_bytes(b"recovery")

    class FakeSealedExporter:
        instances = []

        def __init__(self, *args, cleanup_sources=True, **kwargs):
            self.cleanup_sources = cleanup_sources
            self.errors = []
            self.exported_files = [output]
            self.recovery_files = (recovery,)
            self.cleanup_warnings = (
                f"临时采集缓存清理失败，已保留 {recovery}：injected",
            )
            self.closed = False
            self.__class__.instances.append(self)

        @property
        def queue_ratio(self):
            return 0.0

        def close(self):
            self.closed = True

    monkeypatch.setattr(
        session_manager, "SealedSpoolExporter", FakeSealedExporter
    )
    manager = StorageSessionManager(tmp_path)
    acquisition = request(0)
    closed = []
    warnings = []
    manager.session_closed.connect(
        lambda task_id, files, errors: closed.append((task_id, files, errors))
    )
    manager.warning_event.connect(
        lambda task_id, message: warnings.append((task_id, message))
    )

    manager.start_sealed_export(
        acquisition,
        [ready_device(0)],
        {
            0: {
                "spool_path": str(recovery),
                "persisted_frames": 1,
            }
        },
        cleanup_sources=False,
    )

    assert wait_until(lambda: bool(closed))
    assert FakeSealedExporter.instances[0].cleanup_sources is False
    assert closed == [
        (
            acquisition.task_id,
            [str(output), str(recovery)],
            [],
        )
    ]
    assert len(warnings) == 1
    assert warnings[0][0] == acquisition.task_id
    assert str(recovery) in warnings[0][1]
