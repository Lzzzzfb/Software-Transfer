import csv

from openpyxl import load_workbook
import pytest

from spectrometer.domain.models import SpectrumFrame
from spectrometer.processing.processor import (
    ProcessingSnapshot,
    SpectrumProcessor,
)
from spectrometer.storage.process_worker import export_processed_batch


def test_worker_exports_only_final_calibrated_airpls_values(tmp_path):
    pixel_count = 64
    pixels = [
        1000 + index + (500 if 28 <= index <= 35 else 0)
        for index in range(pixel_count)
    ]
    frame = SpectrumFrame.create(
        0, 7 << 8, pixels, timestamp_ns=123
    )
    csv_path = tmp_path / "result.csv"
    xlsx_path = tmp_path / "result.xlsx"
    snapshot = ProcessingSnapshot(
        mode="dark_subtract",
        wavelengths=tuple(500.0 + index for index in range(pixel_count)),
        intensity_calibration=(2.0,) * pixel_count,
        intensity_calibration_id="cal-003",
        background=(100.0,) * pixel_count,
        baseline_enabled=True,
        baseline_lam=3e5,
        baseline_order=3,
        baseline_max_iter=12,
    )
    expected = SpectrumProcessor().process_frame(frame, snapshot).values

    result = export_processed_batch(
        {0: [frame]},
        processing_snapshots={0: snapshot},
        csv_targets={0: str(csv_path)},
        xlsx_target=str(xlsx_path),
        wavelengths_by_device={0: snapshot.wavelengths},
        device_labels={0: "003"},
        device_metadata={0: {"Serial": "003"}},
        session_metadata={"Session ID": "test"},
    )

    assert result["files"] == [str(csv_path), str(xlsx_path)]
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    header_row = next(
        index
        for index, row in enumerate(rows)
        if row[:2] == ["像素序号", "波长 (nm)"]
    )
    metadata = {
        row[0][2:]: row[1]
        for row in rows[:header_row]
        if len(row) >= 2 and row[0].startswith("# ")
    }
    assert metadata["强度校准 ID"] == "cal-003"
    assert metadata["已应用 airPLS"] == "是"
    assert metadata["airPLS λ"] == "300000.0"
    assert metadata["airPLS 差分阶数"] == "3"
    assert metadata["airPLS 最大迭代次数"] == "12"
    assert metadata["处理管线版本"] == "2"
    csv_values = [
        float(rows[header_row + index + 1][2])
        for index in range(pixel_count)
    ]
    assert csv_values == pytest.approx(expected)

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = workbook["003"]
    xlsx_values = [
        sheet.cell(index + 2, 3).value for index in range(pixel_count)
    ]
    assert xlsx_values == pytest.approx(expected)
    summary = workbook["采集概要"]
    headings = [
        summary.cell(summary.max_row - 1, column).value
        for column in range(1, summary.max_column + 1)
    ]
    assert "处理模式" in headings
    assert "强度校准" in headings
    assert "airPLS" in headings
    workbook.close()
