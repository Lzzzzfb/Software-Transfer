from openpyxl import load_workbook

from spectrometer.domain.enums import StorageFormat
from spectrometer.domain.models import AcquisitionSession, SpectrumFrame
from spectrometer.storage.coordinator import BatchStorageCoordinator


def test_combined_multi_device_coordinator_exports_only_one_workbook(
    tmp_path,
):
    session = AcquisitionSession.create(
        [0, 1], batch_size=2, storage_format=StorageFormat.CSV_EXCEL
    )
    coordinator = BatchStorageCoordinator(
        tmp_path,
        session,
        wavelengths_by_device={0: (400.0, 401.0), 1: (800.0, 801.0)},
        device_labels={0: "设备A", 1: "设备B"},
    )
    for sequence in range(2):
        for device_id in (0, 1):
            coordinator.submit(
                SpectrumFrame.create(device_id, sequence << 8, [sequence, sequence + 1])
            )
    coordinator.close()

    csv_files = list(tmp_path.glob("*.csv"))
    xlsx_files = list(tmp_path.glob("*.xlsx"))
    assert len(csv_files) == 0
    assert len(xlsx_files) == 1
    assert not list(tmp_path.glob("*.part"))
    workbook = load_workbook(xlsx_files[0], read_only=True)
    assert workbook.sheetnames == ["采集概要", "设备A", "设备B"]


def test_combined_single_device_coordinator_exports_only_csv(tmp_path):
    session = AcquisitionSession.create(
        [0], batch_size=2, storage_format=StorageFormat.CSV_EXCEL
    )
    coordinator = BatchStorageCoordinator(
        tmp_path,
        session,
        wavelengths_by_device={0: (400.0, 401.0)},
        device_labels={0: "SN003"},
    )
    coordinator.submit(SpectrumFrame.create(0, 1 << 8, [10, 20]))
    coordinator.submit(SpectrumFrame.create(0, 2 << 8, [11, 21]))
    coordinator.close()

    assert len(list(tmp_path.glob("*.csv"))) == 1
    assert not list(tmp_path.glob("*.xlsx"))
    assert not list(tmp_path.glob("*.part"))


def test_export_failure_exits_storage_thread_and_keeps_recovery_spool(
    tmp_path, monkeypatch
):
    session = AcquisitionSession.create(
        [0], batch_size=2, storage_format=StorageFormat.CSV
    )
    coordinator = BatchStorageCoordinator(tmp_path, session)
    coordinator.submit(SpectrumFrame.create(0, 1 << 8, [10, 20]))

    def fail_export(*args, **kwargs):
        raise OSError("injected export failure")

    monkeypatch.setattr(coordinator, "_export_batch", fail_export)
    coordinator.close(timeout=0.2)

    assert not coordinator._thread.is_alive()
    assert coordinator.errors == ["injected export failure"]
    assert coordinator.spool_path.exists()
