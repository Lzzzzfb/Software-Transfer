import csv

from openpyxl import load_workbook

from spectrometer.domain.models import SpectrumFrame
from spectrometer.storage.csv_exporter import export_device_csv
from spectrometer.storage.xlsx_exporter import export_workbook


def frames(device_id, offset=0):
    return [
        SpectrumFrame.create(device_id, (index + 1) << 8, [offset + index, offset + index + 10])
        for index in range(2)
    ]


def test_csv_is_one_frame_per_column_and_one_pixel_per_row(tmp_path):
    path = export_device_csv(
        tmp_path / "设备一.csv",
        1,
        frames(1),
        wavelengths=(400.0, 401.5),
        metadata={"Serial": "SN001"},
    )
    with path.open(encoding="utf-8-sig", newline="") as file:
        rows = list(csv.reader(file))
    header_index = next(
        index
        for index, row in enumerate(rows)
        if row and row[0] == "像素序号"
    )
    assert rows[0] == ["# ZGCAI 批量光谱数据"]
    assert ["# 生产序列号", "SN001"] in rows[:header_index]
    assert rows[header_index][:2] == ["像素序号", "波长 (nm)"]
    assert rows[header_index][2].startswith("第 0001 帧_序号_")
    assert len(rows[header_index]) == 4
    assert rows[header_index + 1][:2] == ["0", "400.000000"]
    assert rows[header_index + 2][0] == "1"


def test_excel_uses_summary_and_separate_device_sheets(tmp_path):
    path = export_workbook(
        tmp_path / "批量.xlsx",
        {1: frames(1), 2: frames(2, 100)},
        wavelengths_by_device={1: (400.0, 401.0), 2: (900.0, 901.0)},
        device_labels={1: "SN001_可见光", 2: "SN002_近红外"},
        session_metadata={
            "Session ID": "session-1",
            "Sync Mode": "software",
            "Background Applied": True,
        },
        device_metadata={
            1: {
                "Processing Mode": "raw",
                "airPLS Applied": False,
            },
            2: {
                "Processing Mode": "dark_subtract",
                "airPLS Applied": True,
            },
        },
    )
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert workbook.sheetnames == ["采集概要", "SN001_可见光", "SN002_近红外"]
    summary = workbook["采集概要"]
    summary_values = {
        row[0]: row[1]
        for row in summary.iter_rows(values_only=True)
        if row and len(row) >= 2 and row[0]
    }
    assert summary_values["会话编号"] == "session-1"
    assert summary_values["同步方式"] == "软件同步"
    assert summary_values["已应用背景"] == "是"
    summary_rows = list(summary.iter_rows(values_only=True))
    device_header = next(
        index for index, row in enumerate(summary_rows) if row[0] == "设备"
    )
    assert summary_rows[device_header + 1][5] == "原始强度"
    assert summary_rows[device_header + 2][5] == "扣背景"
    visible = workbook["SN001_可见光"]
    infrared = workbook["SN002_近红外"]
    assert visible.cell(1, 1).value == "像素序号"
    assert visible.cell(1, 2).value == "波长 (nm)"
    assert visible.cell(1, 3).value.startswith("第 0001 帧_序号_")
    assert visible.cell(2, 2).value == 400
    assert infrared.cell(2, 2).value == 900
    assert infrared.cell(2, 3).value == 100
